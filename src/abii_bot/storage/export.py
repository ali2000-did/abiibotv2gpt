"""خروجی‌گیری: CSV (سازگار با Excel فارسی)، XLSX (راست‌به‌چپ) و JSON."""
from __future__ import annotations

import csv
import json
import logging
from pathlib import Path

from ..models import Lead

logger = logging.getLogger(__name__)

# ترتیب و نام فارسی ستون‌ها — CSV با utf-8-sig تا اکسل فارسی را درست باز کند
PERSIAN_HEADERS = {
    "source": "پلتفرم",
    "title": "عنوان",
    "phones": "شماره تماس",
    "seller_name": "نام فروشنده",
    "city": "شهر",
    "category": "دسته‌بندی",
    "price": "قیمت",
    "emails": "ایمیل",
    "site": "سایت فروشنده",
    "url": "لینک",
    "description": "توضیحات",
    "quality_score": "امتیاز کیفیت",
    "status": "وضعیت",
    "first_seen": "اولین برداشت",
}
EXPORT_COLUMNS = list(PERSIAN_HEADERS)


def _lead_row(lead: Lead) -> dict:
    d = lead.to_dict()
    return {
        "source": d["source"],
        "title": d["title"] or "",
        "phones": d["phones"],
        "seller_name": d["seller_name"] or "",
        "city": d["city"] or "",
        "category": d["category"] or "",
        "price": d["price"] or "",
        "emails": d["emails"],
        "site": d["site"] or "",
        "url": d["url"],
        "description": d["description"] or "",
        "quality_score": d["quality_score"],
        "status": d["status"],
        "first_seen": d["first_seen"],
    }


def export_leads(
    leads: list[Lead],
    out_dir: Path | str,
    formats: list[str] | None = None,
    name_prefix: str = "leads",
    only_with_contact: bool = False,
) -> dict[str, Path]:
    """تولید فایل‌های خروجی؛ مسیر فایل‌های ساخته‌شده را برمی‌گرداند.

    only_with_contact=True → فقط لیدهای دارای شماره یا ایمیل (پیش‌فرض موتور):
    ردیف‌های بی‌ارزش (بدون هیچ راه تماسی) در اکسل نمی‌آیند.
    """
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    formats = formats or ["csv"]
    paths: dict[str, Path] = {}
    if only_with_contact:
        leads = [l for l in leads if l.phones or l.emails]
    rows = [_lead_row(l) for l in leads]

    if "csv" in formats:
        p = out / f"{name_prefix}.csv"
        with open(p, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=EXPORT_COLUMNS)
            w.writerow(PERSIAN_HEADERS)
            w.writerows(rows)
        paths["csv"] = p

    if "xlsx" in formats:
        p = out / f"{name_prefix}.xlsx"
        _write_xlsx(rows, p)
        paths["xlsx"] = p

    if "json" in formats:
        p = out / f"{name_prefix}.json"
        p.write_text(
            json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        paths["json"] = p

    for fmt, p in paths.items():
        logger.info("exported %s -> %s", fmt, p)
    return paths


def _write_xlsx(rows: list[dict], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "لیدها"
    ws.sheet_view.rightToLeft = True  # شیت راست‌به‌چپ برای فارسی

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    center = Alignment(horizontal="center")

    ws.append([PERSIAN_HEADERS[c] for c in EXPORT_COLUMNS])
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for r in rows:
        ws.append([r[c] for c in EXPORT_COLUMNS])
    for col, width in zip(ws.columns, [10, 40, 18, 20, 12, 18, 16, 24, 30, 46, 40, 12, 12, 20]):
        ws.column_dimensions[col[0].column_letter].width = width

    wb.save(path)
